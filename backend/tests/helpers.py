import io

from docx import Document as DocxDocument
from httpx import AsyncClient


def make_pdf(pages: list[str]) -> bytes:
    """Build a minimal, valid multi-page PDF with one line of Helvetica text per page."""
    page_ids = [4 + 2 * i for i in range(len(pages))]
    objects: dict[int, str] = {
        1: "<< /Type /Catalog /Pages 2 0 R >>",
        2: f"<< /Type /Pages /Kids [{' '.join(f'{p} 0 R' for p in page_ids)}] /Count {len(pages)} >>",
        3: "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    for page_id, text in zip(page_ids, pages, strict=True):
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET"
        objects[page_id] = (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page_id + 1} 0 R >>"
        )
        objects[page_id + 1] = f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream"

    out = b"%PDF-1.4\n"
    offsets: dict[int, int] = {}
    for number in sorted(objects):
        offsets[number] = len(out)
        out += f"{number} 0 obj\n{objects[number]}\nendobj\n".encode("latin-1")
    xref = len(out)
    size = max(objects) + 1
    out += f"xref\n0 {size}\n0000000000 65535 f \n".encode()
    for number in range(1, size):
        out += f"{offsets[number]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def make_docx(heading: str, paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = DocxDocument()
    document.add_heading(heading, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def register(
    client: AsyncClient, email: str, password: str = "correct-horse-battery"
) -> dict[str, str]:
    response = await client.post("/auth/register", json={"email": email, "password": password})
    assert response.status_code == 201, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


async def upload(
    client: AsyncClient,
    headers: dict[str, str],
    filename: str,
    content: bytes,
    *,
    tags: str = "",
    visibility: str = "private",
    expect: int = 202,
    calibrate: bool = True,
) -> dict:
    response = await client.post(
        "/documents/upload",
        headers=headers,
        files={"file": (filename, content)},
        data={"tags": tags, "visibility": visibility},
    )
    assert response.status_code == expect, response.text
    if calibrate and expect == 202:
        # ASGITransport runs background tasks before returning, so processing is done.
        calibrated = await client.post("/documents/calibrate", headers=headers)
        assert calibrated.status_code == 202, calibrated.text
    return response.json()


async def ask(client: AsyncClient, headers: dict[str, str], question: str, **extra) -> dict:
    response = await client.post("/chat", headers=headers, json={"question": question, **extra})
    assert response.status_code == 200, response.text
    return response.json()
